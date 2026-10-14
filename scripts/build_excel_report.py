"""
Week 4: Capstone Project — Executive Financial Analysis Excel Report Generator
Generates: Week4_Financial_Analysis_Report.xlsx with Executive Dashboard, Pivot Summaries & Raw Data
"""

import os
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

def create_excel_report():
    print("=" * 80)
    print("WEEK 4: CAPSTONE PROJECT — EXCEL FINANCIAL ANALYSIS REPORT GENERATOR")
    print("=" * 80)
    
    script_dir = os.path.dirname(os.path.abspath(__file__))
    data_path = os.path.join(script_dir, "..", "data", "financial_sample.xlsx")
    report_output = os.path.join(script_dir, "..", "reports", "Week4_Financial_Analysis_Report.xlsx")
    os.makedirs(os.path.dirname(report_output), exist_ok=True)
    
    df = pd.read_excel(data_path)
    df.columns = df.columns.str.strip()
    
    print("Writing base workbook with cleaned financial data...")
    with pd.ExcelWriter(report_output, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Financial_Data", index=False)
        
    wb = openpyxl.load_workbook(report_output)
    
    # Styles definition
    navy_fill = PatternFill(start_color="0F2537", end_color="0F2537", fill_type="solid")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    accent_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    sub_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    kpi_card_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
    
    font_main_title = Font(name="Calibri", size=18, bold=True, color="FFFFFF")
    font_sub_title = Font(name="Calibri", size=11, italic=True, color="D9E1F2")
    font_section = Font(name="Calibri", size=13, bold=True, color="0F2537")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_kpi_val = Font(name="Calibri", size=18, bold=True, color="0F2537")
    font_kpi_sub = Font(name="Calibri", size=9, italic=True, color="2E75B6")
    font_bold = Font(name="Calibri", size=11, bold=True)
    font_regular = Font(name="Calibri", size=11)
    
    thin_border = Border(left=Side(style="thin", color="D9D9D9"),
                         right=Side(style="thin", color="D9D9D9"),
                         top=Side(style="thin", color="D9D9D9"),
                         bottom=Side(style="thin", color="D9D9D9"))
    double_bottom = Border(top=Side(style="thin", color="D9D9D9"),
                           bottom=Side(style="double", color="1F4E79"))
                           
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    
    # Format Financial_Data sheet headers
    ws_data = wb["Financial_Data"]
    ws_data.views.sheetView[0].showGridLines = True
    for col_idx in range(1, len(df.columns) + 1):
        c = ws_data.cell(row=1, column=col_idx)
        c.fill = header_fill
        c.font = font_header
        c.alignment = align_center

    # -------------------------------------------------------------------------
    # Sheet 1: Executive_Dashboard
    # -------------------------------------------------------------------------
    print("Creating Executive_Dashboard sheet...")
    ws_dash = wb.create_sheet(title="Executive_Dashboard", index=0)
    ws_dash.views.sheetView[0].showGridLines = True
    
    # Title Banner
    for r in range(2, 4):
        for col_letter in ["B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M"]:
            ws_dash[f"{col_letter}{r}"].fill = navy_fill
            
    ws_dash["B2"] = "EXECUTIVE SALES & PROFITABILITY DASHBOARD"
    ws_dash["B2"].font = font_main_title
    ws_dash["B3"] = "Week 4 Capstone Project: End-to-End Enterprise Financial Analytics (Microsoft Financial Sample)"
    ws_dash["B3"].font = font_sub_title
    
    # KPI Cards (Row 5 to 7)
    kpis = [
        ("TOTAL NET SALES", "$118,726,350", "700 Transactions", "B", "C", header_fill),
        ("TOTAL NET PROFIT", "$16,893,702", "Healthy Cashflow", "D", "E", accent_fill),
        ("PROFIT MARGIN", "14.23%", "Target > 12.0%", "F", "G", header_fill),
        ("UNITS SOLD", "1,125,806", "Avg 1,608 / Order", "H", "I", accent_fill),
        ("TOP SEGMENT", "Government", "$52.5M (44.2% Share)", "J", "K", header_fill),
        ("TOP PRODUCT", "Paseo", "$33.0M Revenue", "L", "M", accent_fill),
    ]
    
    for title, val, sub, c1, c2, b_fill in kpis:
        ws_dash[f"{c1}5"] = title
        ws_dash[f"{c1}5"].fill = b_fill
        ws_dash[f"{c1}5"].font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
        ws_dash[f"{c1}5"].alignment = align_center
        ws_dash[f"{c2}5"].fill = b_fill
        
        ws_dash[f"{c1}6"] = val
        ws_dash[f"{c1}6"].font = font_kpi_val
        ws_dash[f"{c1}6"].alignment = align_center
        ws_dash[f"{c1}6"].fill = kpi_card_fill
        ws_dash[f"{c2}6"].fill = kpi_card_fill
        
        ws_dash[f"{c1}7"] = sub
        ws_dash[f"{c1}7"].font = font_kpi_sub
        ws_dash[f"{c1}7"].alignment = align_center
        ws_dash[f"{c1}7"].fill = kpi_card_fill
        ws_dash[f"{c2}7"].fill = kpi_card_fill
        
        for r in range(5, 8):
            for col in [c1, c2]:
                ws_dash[f"{col}{r}"].border = thin_border
                
    # Table 1: Country Breakdown (Row 9 to 16, Cols B to G)
    ws_dash["B9"] = "1. Financial Performance by Country"
    ws_dash["B9"].font = font_section
    
    country_headers = ["Country", "Units Sold", "Net Sales ($)", "% Sales Share", "Net Profit ($)", "Margin %"]
    for i, h in enumerate(country_headers, start=2):
        c = ws_dash.cell(row=10, column=i, value=h)
        c.fill = header_fill
        c.font = font_header
        c.alignment = align_center
        c.border = thin_border
        
    country_summary = df.groupby('Country').agg(
        Units=('Units Sold', 'sum'),
        Sales=('Sales', 'sum'),
        Profit=('Profit', 'sum')
    ).sort_values(by='Sales', ascending=False).reset_index()
    
    for idx, row in country_summary.iterrows():
        r = 11 + idx
        ws_dash.cell(row=r, column=2, value=row["Country"]).alignment = align_left
        ws_dash.cell(row=r, column=3, value=row["Units"]).alignment = align_right
        ws_dash.cell(row=r, column=4, value=row["Sales"]).alignment = align_right
        ws_dash.cell(row=r, column=5, value=f"=D{r}/$D$16").alignment = align_right
        ws_dash.cell(row=r, column=6, value=row["Profit"]).alignment = align_right
        ws_dash.cell(row=r, column=7, value=f"=F{r}/D{r}").alignment = align_right
        
        ws_dash.cell(row=r, column=3).number_format = "#,##0"
        ws_dash.cell(row=r, column=4).number_format = "$#,##0.00"
        ws_dash.cell(row=r, column=5).number_format = "0.00%"
        ws_dash.cell(row=r, column=6).number_format = "$#,##0.00"
        ws_dash.cell(row=r, column=7).number_format = "0.00%"
        
        for col_idx in range(2, 8):
            ws_dash.cell(row=r, column=col_idx).border = thin_border
            ws_dash.cell(row=r, column=col_idx).font = font_regular
            
    # Country Total Row
    ws_dash.cell(row=16, column=2, value="Grand Total").font = font_bold
    ws_dash.cell(row=16, column=3, value="=SUM(C11:C15)").font = font_bold
    ws_dash.cell(row=16, column=4, value="=SUM(D11:D15)").font = font_bold
    ws_dash.cell(row=16, column=5, value="=SUM(E11:E15)").font = font_bold
    ws_dash.cell(row=16, column=6, value="=SUM(F11:F15)").font = font_bold
    ws_dash.cell(row=16, column=7, value="=F16/D16").font = font_bold
    
    ws_dash.cell(row=16, column=3).number_format = "#,##0"
    ws_dash.cell(row=16, column=4).number_format = "$#,##0.00"
    ws_dash.cell(row=16, column=5).number_format = "0.00%"
    ws_dash.cell(row=16, column=6).number_format = "$#,##0.00"
    ws_dash.cell(row=16, column=7).number_format = "0.00%"
    
    for c in range(2, 8):
        ws_dash.cell(row=16, column=c).border = double_bottom
        ws_dash.cell(row=16, column=c).fill = sub_fill
        
    # Table 2: Segment Breakdown (Row 9 to 16, Cols I to M)
    ws_dash["I9"] = "2. Financial Performance by Market Segment"
    ws_dash["I9"].font = font_section
    
    seg_headers = ["Segment", "Net Sales ($)", "% Sales", "Net Profit ($)", "Margin %"]
    for i, h in enumerate(seg_headers, start=9):
        c = ws_dash.cell(row=10, column=i, value=h)
        c.fill = header_fill
        c.font = font_header
        c.alignment = align_center
        c.border = thin_border
        
    seg_summary = df.groupby('Segment').agg(
        Sales=('Sales', 'sum'),
        Profit=('Profit', 'sum')
    ).sort_values(by='Sales', ascending=False).reset_index()
    
    for idx, row in seg_summary.iterrows():
        r = 11 + idx
        ws_dash.cell(row=r, column=9, value=row["Segment"]).alignment = align_left
        ws_dash.cell(row=r, column=10, value=row["Sales"]).alignment = align_right
        ws_dash.cell(row=r, column=11, value=f"=J{r}/$J$16").alignment = align_right
        ws_dash.cell(row=r, column=12, value=row["Profit"]).alignment = align_right
        ws_dash.cell(row=r, column=13, value=f"=L{r}/J{r}").alignment = align_right
        
        ws_dash.cell(row=r, column=10).number_format = "$#,##0.00"
        ws_dash.cell(row=r, column=11).number_format = "0.00%"
        ws_dash.cell(row=r, column=12).number_format = "$#,##0.00"
        ws_dash.cell(row=r, column=13).number_format = "0.00%"
        
        for col_idx in range(9, 14):
            ws_dash.cell(row=r, column=col_idx).border = thin_border
            ws_dash.cell(row=r, column=col_idx).font = font_regular
            
    # Segment Total Row
    ws_dash.cell(row=16, column=9, value="Grand Total").font = font_bold
    ws_dash.cell(row=16, column=10, value="=SUM(J11:J15)").font = font_bold
    ws_dash.cell(row=16, column=11, value="=SUM(K11:K15)").font = font_bold
    ws_dash.cell(row=16, column=12, value="=SUM(L11:L15)").font = font_bold
    ws_dash.cell(row=16, column=13, value="=L16/J16").font = font_bold
    
    ws_dash.cell(row=16, column=10).number_format = "$#,##0.00"
    ws_dash.cell(row=16, column=11).number_format = "0.00%"
    ws_dash.cell(row=16, column=12).number_format = "$#,##0.00"
    ws_dash.cell(row=16, column=13).number_format = "0.00%"
    
    for c in range(9, 14):
        ws_dash.cell(row=16, column=c).border = double_bottom
        ws_dash.cell(row=16, column=c).fill = sub_fill
        
    # Table 3: Product Breakdown (Row 18 to 25, Cols B to G)
    ws_dash["B18"] = "3. Product Family Profitability"
    ws_dash["B18"].font = font_section
    
    prod_headers = ["Product", "Units Sold", "Net Sales ($)", "% Sales", "Net Profit ($)", "Margin %"]
    for i, h in enumerate(prod_headers, start=2):
        c = ws_dash.cell(row=19, column=i, value=h)
        c.fill = header_fill
        c.font = font_header
        c.alignment = align_center
        c.border = thin_border
        
    prod_summary = df.groupby('Product').agg(
        Units=('Units Sold', 'sum'),
        Sales=('Sales', 'sum'),
        Profit=('Profit', 'sum')
    ).sort_values(by='Sales', ascending=False).reset_index()
    
    for idx, row in prod_summary.iterrows():
        r = 20 + idx
        ws_dash.cell(row=r, column=2, value=row["Product"]).alignment = align_left
        ws_dash.cell(row=r, column=3, value=row["Units"]).alignment = align_right
        ws_dash.cell(row=r, column=4, value=row["Sales"]).alignment = align_right
        ws_dash.cell(row=r, column=5, value=f"=D{r}/$D$26").alignment = align_right
        ws_dash.cell(row=r, column=6, value=row["Profit"]).alignment = align_right
        ws_dash.cell(row=r, column=7, value=f"=F{r}/D{r}").alignment = align_right
        
        ws_dash.cell(row=r, column=3).number_format = "#,##0"
        ws_dash.cell(row=r, column=4).number_format = "$#,##0.00"
        ws_dash.cell(row=r, column=5).number_format = "0.00%"
        ws_dash.cell(row=r, column=6).number_format = "$#,##0.00"
        ws_dash.cell(row=r, column=7).number_format = "0.00%"
        
        for col_idx in range(2, 8):
            ws_dash.cell(row=r, column=col_idx).border = thin_border
            ws_dash.cell(row=r, column=col_idx).font = font_regular
            
    # Product Total Row
    ws_dash.cell(row=26, column=2, value="Grand Total").font = font_bold
    ws_dash.cell(row=26, column=3, value="=SUM(C20:C25)").font = font_bold
    ws_dash.cell(row=26, column=4, value="=SUM(D20:D25)").font = font_bold
    ws_dash.cell(row=26, column=5, value="=SUM(E20:E25)").font = font_bold
    ws_dash.cell(row=26, column=6, value="=SUM(F20:F25)").font = font_bold
    ws_dash.cell(row=26, column=7, value="=F26/D26").font = font_bold
    
    ws_dash.cell(row=26, column=3).number_format = "#,##0"
    ws_dash.cell(row=26, column=4).number_format = "$#,##0.00"
    ws_dash.cell(row=26, column=5).number_format = "0.00%"
    ws_dash.cell(row=26, column=6).number_format = "$#,##0.00"
    ws_dash.cell(row=26, column=7).number_format = "0.00%"
    
    for c in range(2, 8):
        ws_dash.cell(row=26, column=c).border = double_bottom
        ws_dash.cell(row=26, column=c).fill = sub_fill
        
    # Table 4: Discount Band Impact Analysis (Row 18 to 23, Cols I to M)
    ws_dash["I18"] = "4. Impact of Discount Bands"
    ws_dash["I18"].font = font_section
    
    disc_headers = ["Discount Band", "Orders", "Net Sales ($)", "Net Profit ($)", "Margin %"]
    for i, h in enumerate(disc_headers, start=9):
        c = ws_dash.cell(row=19, column=i, value=h)
        c.fill = header_fill
        c.font = font_header
        c.alignment = align_center
        c.border = thin_border
        
    disc_summary = df.groupby('Discount Band').agg(
        Orders=('Sales', 'count'),
        Sales=('Sales', 'sum'),
        Profit=('Profit', 'sum')
    ).reindex(['None', 'Low', 'Medium', 'High']).reset_index()
    
    for idx, row in disc_summary.iterrows():
        r = 20 + idx
        ws_dash.cell(row=r, column=9, value=row["Discount Band"]).alignment = align_left
        ws_dash.cell(row=r, column=10, value=row["Orders"]).alignment = align_right
        ws_dash.cell(row=r, column=11, value=row["Sales"]).alignment = align_right
        ws_dash.cell(row=r, column=12, value=row["Profit"]).alignment = align_right
        ws_dash.cell(row=r, column=13, value=f"=L{r}/K{r}").alignment = align_right
        
        ws_dash.cell(row=r, column=10).number_format = "#,##0"
        ws_dash.cell(row=r, column=11).number_format = "$#,##0.00"
        ws_dash.cell(row=r, column=12).number_format = "$#,##0.00"
        ws_dash.cell(row=r, column=13).number_format = "0.00%"
        
        for col_idx in range(9, 14):
            ws_dash.cell(row=r, column=col_idx).border = thin_border
            ws_dash.cell(row=r, column=col_idx).font = font_regular

    # Add Chart to Dashboard
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Country Net Sales & Profit ($ USD)"
    chart.y_axis.title = "USD ($)"
    chart.x_axis.title = "Country"
    chart.height = 11
    chart.width = 17
    
    sales_ref = Reference(ws_dash, min_col=4, min_row=10, max_row=15)
    profit_ref = Reference(ws_dash, min_col=6, min_row=10, max_row=15)
    cats_ref = Reference(ws_dash, min_col=2, min_row=11, max_row=15)
    chart.add_data(sales_ref, titles_from_data=True)
    chart.add_data(profit_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    
    ws_dash.add_chart(chart, "B28")

    # Key Insights Section
    ins_row = 49
    ws_dash[f"B{ins_row}"] = "Strategic Executive Takeaways & Recommendations:"
    ws_dash[f"B{ins_row}"].font = font_section
    
    recs = [
        "1. Core Revenue Pillar: The Government segment drives $52.50M (44.22% of total sales) with healthy 21.69% margins. Sustaining long-term municipal contracts must remain top strategic priority.",
        "2. High Volume / Margin Squeeze in Small Business: Small Business accounts for $42.43M in sales (35.74%) but generates only 9.77% net margin due to aggressive discounting.",
        "3. Discount Band Danger: High discount bands compress profit margins from 19.5% (None/Low) down to 7.8% (High). Implement discount threshold governance to avoid value destruction.",
        "4. Balanced Global Footprint: USA ($25.03M), Canada ($24.89M), France ($24.35M), and Germany ($23.51M) exhibit remarkably balanced demand. Cross-selling Paseo and VTT across European channels offers strong upside."
    ]
    for idx, text in enumerate(recs, start=ins_row+1):
        ws_dash[f"B{idx}"] = text
        ws_dash[f"B{idx}"].font = font_regular

    # Adjust Column Widths
    ws_dash.column_dimensions["A"].width = 3
    ws_dash.column_dimensions["B"].width = 16
    ws_dash.column_dimensions["C"].width = 16
    ws_dash.column_dimensions["D"].width = 20
    ws_dash.column_dimensions["E"].width = 16
    ws_dash.column_dimensions["F"].width = 20
    ws_dash.column_dimensions["G"].width = 14
    ws_dash.column_dimensions["H"].width = 4
    ws_dash.column_dimensions["I"].width = 22
    ws_dash.column_dimensions["J"].width = 18
    ws_dash.column_dimensions["K"].width = 16
    ws_dash.column_dimensions["L"].width = 18
    ws_dash.column_dimensions["M"].width = 14
    
    wb.save(report_output)
    print(f"Executive Financial Report saved successfully to: {report_output}")

if __name__ == "__main__":
    create_excel_report()
